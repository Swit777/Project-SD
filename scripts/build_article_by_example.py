from __future__ import annotations

import argparse
from copy import deepcopy
import json
from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from jinja2 import Environment, StrictUndefined

import build_game_concept_article as article
from article_retrieval_figures import make_figures

ROOT = article.ROOT
REFERENCE = ROOT / "Статья_цены_подержанных_автомобилей_.docx"
TEMPLATE = ROOT / "docs/game_concept_article_short_template_ru.md"
OUTPUT = ROOT / "reports/game_concept/article_by_example"
DOCX = ROOT / "Статья_прогнозирование_игровых_концепций.docx"
TITLE_RO = ("PROGNOZAREA POTENȚIALULUI COMERCIAL AL CONCEPTELOR DE JOCURI "
            "PE BAZA GENURILOR ȘI A COMBINAȚIILOR DE MECANICI DE JOC, "
            "FOLOSIND DATELE PLATFORMEI STEAM")
TITLE_RU = ("ПРОГНОЗИРОВАНИЕ КОММЕРЧЕСКОГО ПОТЕНЦИАЛА ИГРОВЫХ КОНЦЕПЦИЙ "
            "ПО ЖАНРУ И СОЧЕТАНИЯМ ИГРОВЫХ МЕХАНИК НА ДАННЫХ STEAM")
FIGURES = ["01_dataset.png", "02_pipeline.png", "03_models.png", "04_hypothesis.png",
           "06_importance.png", "05_calibration.png", "08_application.png"]


def use_paragraph_format(paragraph, sample):
    existing = paragraph._p.pPr
    if existing is not None:
        paragraph._p.remove(existing)
    if sample._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(sample._p.pPr))
    paragraph.paragraph_format.widow_control = True


def set_fonts(paragraph, size=12):
    for run in paragraph._p.xpath(".//w:r"):
        properties = run.get_or_add_rPr()
        fonts = properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        for key in ["ascii", "hAnsi", "cs", "eastAsia"]:
            fonts.set(qn(f"w:{key}"), "Times New Roman")
        for key in ["asciiTheme", "hAnsiTheme", "csTheme", "eastAsiaTheme"]:
            fonts.attrib.pop(qn(f"w:{key}"), None)
        for name in ["sz", "szCs"]:
            element = properties.find(qn(f"w:{name}"))
            if element is None:
                element = OxmlElement(f"w:{name}")
                properties.append(element)
            element.set(qn("w:val"), str(round(size * 2)))
        for color in properties.findall(qn("w:color")):
            color.set(qn("w:val"), "000000")


def paragraph_from_sample(document, sample, text, size=12, bold=False, italic=False):
    paragraph = document.add_paragraph()
    use_paragraph_format(paragraph, sample)
    article.append_inline(paragraph, text)
    set_fonts(paragraph, size)
    for run in paragraph.runs:
        if bold:
            run.bold = True
        if italic:
            run.italic = True
    return paragraph


def prepared_template(reference):
    document = Document(reference)
    for element in list(document.element.body):
        if element.tag != qn("w:sectPr"):
            document.element.body.remove(element)
    # Preserve the template's header, footer and styles, without the source article's body assets.
    for rel_id, relation in list(document.part.rels.items()):
        if relation.reltype.rsplit("/", 1)[-1] in {"image", "hyperlink", "comments", "footnotes", "endnotes", "customXml"}:
            document.part.drop_rel(rel_id)
    for element in list(document.settings.element):
        if element.tag in {qn("w:footnotePr"), qn("w:endnotePr")}:
            document.settings.element.remove(element)
    properties = document.core_properties
    properties.title, properties.subject = TITLE_RU, "Игровая аналитика на данных Steam"
    properties.author = properties.last_modified_by = properties.comments = ""
    properties.keywords = "Steam, игровые концепции, машинное обучение, неопределённость"
    properties.category = "Исследовательская статья"
    footer = document.sections[0].footer
    for paragraph in footer.paragraphs:
        if "Chisinau" in paragraph.text:
            for node in paragraph._p.xpath(".//w:t"):
                node.text = "Chisinau, 09.10.2026" if "Chisinau" in (node.text or "") else ""
    return document


def cover(document, sample, args):
    paragraph_from_sample(document, sample.paragraphs[0], TITLE_RO, size=14, bold=True)
    paragraph_from_sample(document, sample.paragraphs[1], TITLE_RU, size=14, bold=True)
    paragraph_from_sample(document, sample.paragraphs[2], args.author or "[Prenume NUME]", bold=True)
    affiliation = args.affiliation or ("Departamentul Informatică și Ingineria Sistemelor, grupa " + args.group +
        ", Facultatea FCIM, Universitatea Tehnică a Moldovei, Chișinău, Moldova")
    paragraph_from_sample(document, sample.paragraphs[3], affiliation, italic=True)
    paragraph_from_sample(document, sample.paragraphs[4],
                          f"Autorul corespondent: {args.author or '[Prenume NUME]'}, e-mail: {args.email or '[e-mail]'}")
    paragraph_from_sample(document, sample.paragraphs[5],
                          "Îndrumătorul/coordonatorul științific: " + (args.supervisor or "[Nume, titlu științific]"), bold=True)
    document.core_properties.author = args.author


def add_table(document, sample, rows):
    count = len(rows[0])
    width = document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin
    weights = [.27, .73] if count == 2 else [.47, .24, .29] if count == 3 else [.31] + [.69 / (count - 1)] * (count - 1)
    table = document.add_table(rows=0, cols=count)
    table._tbl.replace(table._tbl.tblPr, deepcopy(sample.tables[0]._tbl.tblPr))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, share in zip(table.columns, weights):
        column.width = int(width * share)
    for index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for column, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = int(width * weights[column])
            paragraph = cell.paragraphs[0]
            use_paragraph_format(paragraph, sample.tables[0].rows[min(index, 1)].cells[min(column, 2)].paragraphs[0])
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = index < len(rows) - 1
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column == 0 or count == 2 else WD_ALIGN_PARAGRAPH.CENTER
            article.append_inline(paragraph, text)
            set_fonts(paragraph, 10)
            for run in paragraph.runs:
                run.bold = index == 0
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 1
    spacer.add_run().font.size = Pt(2)


def render_docx(markdown, args):
    sample = Document(REFERENCE)
    document = prepared_template(REFERENCE)
    cover(document, sample, args)
    blocks = article.parse_markdown(markdown, OUTPUT)
    in_results, in_bibliography = False, False
    for block in blocks:
        if block.kind == "heading":
            text = article.plain(block.html)
            if block.level == 2:
                in_results = text == "Результаты и обсуждение"
                in_bibliography = text == "Библиография"
                index = {"Резюме": 6, "Введение": 11, "Библиография": 101}.get(text, 16)
                paragraph = paragraph_from_sample(document, sample.paragraphs[index], block.html, size=14, bold=True)
                paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.page_break_before = text == "Резюме"
            else:
                paragraph = paragraph_from_sample(document, sample.paragraphs[54 if in_results else 17], block.html, bold=True)
                paragraph.paragraph_format.keep_with_next = True
        elif block.kind in {"paragraph", "list"}:
            if in_bibliography:
                paragraph = paragraph_from_sample(document, sample.paragraphs[102], block.html)
            elif article.plain(block.html).startswith("Ключевые слова:"):
                paragraph = paragraph_from_sample(document, sample.paragraphs[10], block.html)
            else:
                paragraph = paragraph_from_sample(document, sample.paragraphs[7], block.html)
        elif block.kind == "caption":
            text = article.plain(block.html)
            label, title = text.split(". ", 1)
            paragraph_from_sample(document, sample.paragraphs[22], label)
            paragraph_from_sample(document, sample.paragraphs[23], title, size=11)
        elif block.kind == "table":
            add_table(document, sample, block.rows)
        elif block.kind == "figure":
            paragraph = paragraph_from_sample(document, sample.paragraphs[32], "")
            paragraph.paragraph_format.keep_with_next = True
            width, _ = article.image_dimensions(block.path, 15.875, 12.6)
            paragraph.add_run().add_picture(str(block.path), width=Cm(width))
            caption = paragraph_from_sample(document, sample.paragraphs[33], block.caption, size=11)
            caption.paragraph_format.keep_together = True
        else:
            raise ValueError(f"Unsupported article block: {block.kind}")
    document.save(DOCX)
    return blocks


def context_for_example(results, test_count):
    context = article.build_context(results, "", "", test_count)
    context["verified_tests"] = test_count
    a = results["audit"]
    context["sampling_table"] = article.markdown_table(["Этап", "Исключено / не выбрано", "Осталось"], [
        ["Строки закреплённого CSV", "", article.integer(a["source_rows"])],
        ["Совокупность содержательных фильтров", article.integer(a["source_rows"] - a["eligible_rows"]), article.integer(a["eligible_rows"])],
        ["Воспроизводимая подвыборка по SHA-256", article.integer(a["eligible_rows"] - a["prepared_games"]), article.integer(a["prepared_games"])],
    ])
    context["features_table"] = article.markdown_table(["Группа", "Содержание"], [
        ["Контекст (6)", "log(1 + цена), log(1 + возраст в годах), число языков; индикаторы Windows, macOS, Linux"],
        ["Жанры (10)", "Бинарные индикаторы поддерживаемых жанров Steam"],
        ["Элементы (36)", "Свидетельства игровых механик, режимов и структуры мира"],
        ["Производный признак (1)", "Количество найденных положительных элементов"],
        ["Основная цель", "Четыре диапазона оценочного числа владельцев SteamSpy"],
        ["Дополнительная цель", "Медианное накопленное время игры; log(1 + часы) при обучении"],
    ])
    predictions = results["tables"]["test_predictions"]
    selected = results["manifest"]["selected_model"]
    predicted = predictions[[f"{selected}_p{k}" for k in range(4)]].to_numpy().argmax(axis=1)
    target = predictions.owners_band.to_numpy()
    labels = ["0–20 тыс.", "20–100 тыс.", "100–500 тыс.", "500 тыс. и более"]
    rows = [[label, int((target == k).sum()), int(((target == k) & (predicted == k)).sum()),
             article.number((predicted[target == k] == k).mean() * 100, 1)] for k, label in enumerate(labels)]
    context["class_table"] = article.markdown_table(["Истинный диапазон", "Игр", "Распознано", "Полнота, %"], rows)
    context["selected_accuracy"] = article.number((target == predicted).mean() * 100, 1)
    temporal = results["tables"]["temporal_backtest"].set_index("model")
    for name, model in [("additive", "hgb_additive"), ("interactions", "hgb_interactions")]:
        context[f"temporal_{name}_loss"] = article.number(temporal.loc[model, "group_log_loss"])
        context[f"temporal_{name}_balanced"] = article.number(temporal.loc[model, "balanced_accuracy"])
    example = results["example"]
    for key in ["20k", "100k", "500k"]:
        context[f"scenario_{key}"] = article.number(example[f"p_at_least_{key}"] * 100, 1)
    context["scenario_hours"] = article.number(example["playtime"]["median_hours_estimate"], 1)
    context["scenario_hours_low"] = article.number(example["playtime"]["low_hours"], 1)
    context["scenario_hours_high"] = article.number(example["playtime"]["high_hours"], 1)
    existing = (article.OUTPUT / "article_ru.md").read_text(encoding="utf-8")
    bibliography = existing.split("## Список источников\n", 1)[1].split("## Приложение", 1)[0]
    entries = []
    for line in bibliography.splitlines():
        match = re.match(r"(\d+)\. (.+)", line)
        if match:
            # Bracketed references and a hanging indent match the supplied sample.
            entries.append(f"[{match.group(1)}]\t{match.group(2)}")
    if len(entries) != 13:
        raise ValueError("Expected the previously verified bibliography of 13 sources")
    entries += ["[14]\tValve. IUserReviewsService: GetAppReviews. [Steamworks Documentation](https://partner.steamgames.com/doc/webapi/IUserReviewsService). Дата обращения: 09.10.2026.",
                "[15]\tscikit-learn developers. TfidfVectorizer. [Документация](https://scikit-learn.org/stable/modules/generated/sklearn.feature_extraction.text.TfidfVectorizer.html). Дата обращения: 09.10.2026."]
    context["bibliography"] = "\n\n".join(entries)
    upgrade = article.REPORT / "analogue_upgrade"
    comparison = json.loads((upgrade / "comparison.json").read_text(encoding="utf-8"))
    for name, digest in comparison["retrieval_code_sha256"].items():
        if article.sha256_file(ROOT / "src/game_concept" / name) != digest:
            raise ValueError("Retrieval comparison is stale: " + name)
    context["analogue_catalog_count"] = article.integer(comparison["catalog_games"])
    case = comparison["cases"]["rpg_coop_crafting"]
    context["analogue_modes_table"] = article.markdown_table(["Режим", "Кандидатов", "Медиана отзывов в top-12", "Покрытие выбранных механик, %"], [
        [title, article.integer(case[key]["eligible_games"]) if key else "4 000 до ранжирования", article.number(case[mode]["median_reviews"], 1),
         article.number(case[mode]["mean_selected_coverage"] * 100, 1)]
        for mode, key, title in [("old", None, "Исходный Jaccard (train)"), ("new", "diagnostics", "Строгий"),
                                  ("balanced", "balanced_diagnostics", "Ближайшие"), ("broad", "broad_diagnostics", "Рыночные референсы")]])
    context["strict_candidates"] = article.integer(case["diagnostics"]["eligible_games"])
    context["balanced_candidates"] = article.integer(case["balanced_diagnostics"]["eligible_games"])
    context["broad_candidates"] = article.integer(case["broad_diagnostics"]["eligible_games"])
    context["old_low_reviews"] = case["old"]["below_30_reviews"]
    context["old_mode_misses"] = case["old"]["unconfirmed_multiplayer"]
    live = json.loads((ROOT / "data/processed/game_concept/review_totals/2686630.json").read_text(encoding="utf-8"))
    context["live_review_count"] = article.integer(live["total_reviews"])
    context["live_review_date"] = live["captured_at_utc"][:10]
    return context, rows


def main():
    parser = argparse.ArgumentParser(description="Build a DOCX using the user's supplied article as its formatting template")
    parser.add_argument("--author", default="")
    parser.add_argument("--email", default="")
    parser.add_argument("--group", default="[grupa]")
    parser.add_argument("--affiliation", default="")
    parser.add_argument("--supervisor", default="")
    args = parser.parse_args()
    results = article.load_results(allow_retrieval_update=True)
    previous_manifest = json.loads((article.OUTPUT / "article_manifest.json").read_text(encoding="utf-8"))
    if article.sha256_file(article.OUTPUT / "article_ru.md") != previous_manifest["outputs_sha256"]["article_ru.md"]:
        raise ValueError("The source article changed; verify its bibliography before adapting it")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    (OUTPUT / "figures").mkdir(exist_ok=True)
    if DOCX.exists():
        previous = OUTPUT / "versions" / (article.sha256_file(DOCX)[:12] + ".docx")
        previous.parent.mkdir(exist_ok=True)
        if not previous.exists():
            shutil.copy2(DOCX, previous)
    for name in FIGURES:
        source = article.OUTPUT / "figures" / name
        if article.sha256_file(source) != previous_manifest["image_sha256"][name]:
            raise ValueError(f"A verified figure changed: {name}")
        shutil.copy2(source, OUTPUT / "figures" / name)
    make_figures(OUTPUT / "figures")
    test_log = (article.REPORT / "analogue_upgrade/tests.log").read_bytes()
    test_text = test_log.decode("utf-16" if test_log.startswith(b"\xff\xfe") else "utf-8")
    test_counts = re.findall(r"Ran (\d+) tests", test_text)
    if not test_counts or not re.search(r"\nOK\s*$", test_text):
        raise ValueError("A completed successful test run is required for the updated article")
    test_count = int(test_counts[-1])
    context, class_rows = context_for_example(results, test_count)
    environment = Environment(undefined=StrictUndefined, autoescape=False, keep_trailing_newline=True)
    markdown = environment.from_string(TEMPLATE.read_text(encoding="utf-8")).render(**context)
    (OUTPUT / "article_body_ru.md").write_text(markdown, encoding="utf-8")
    blocks = render_docx(markdown, args)
    import pandas as pd
    pd.DataFrame(class_rows, columns=["band", "games", "correct", "recall_percent_display"]).to_csv(OUTPUT / "class_diagnostics.csv", index=False)
    manifest = {"article_date": "2026-10-09", "formatting_reference": REFERENCE.name, "reference_sha256": article.sha256_file(REFERENCE),
                "output_docx": DOCX.name, "docx_sha256": article.sha256_file(DOCX),
                "article_body_sha256": article.sha256_file(OUTPUT / "article_body_ru.md"),
                "experiment_input_sha256": results["manifest"]["input_sha256"], "metadata": vars(args),
                "tables": sum(block.kind == "table" for block in blocks), "figures": sum(block.kind == "figure" for block in blocks),
                "sources": 15, "additional_calculation": "Per-class recall and separate retrieval benchmarks; no refitting of the market model",
                "retrieval_code_changes": results["retrieval_code_changes"],
                "retrieval_comparison_sha256": article.sha256_file(article.REPORT / "analogue_upgrade/comparison.json"),
                "test_log_sha256": article.sha256_file(article.REPORT / "analogue_upgrade/tests.log"),
                "figures_sha256": {p.name: article.sha256_file(p) for p in (OUTPUT / "figures").glob("*.png")},
                "source_code_sha256": {path.name: article.sha256_file(path) for path in [Path(__file__), TEMPLATE]},
                "reviewed_tests": test_count}
    (OUTPUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"file": str(DOCX), "tables": manifest["tables"], "figures": manifest["figures"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
