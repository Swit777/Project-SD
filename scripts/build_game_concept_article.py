from __future__ import annotations

import argparse
import importlib.metadata
import json
import shutil
import sys
from dataclasses import dataclass, field
from pathlib import Path

import joblib
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, Rectangle
import numpy as np
import pandas as pd
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from jinja2 import Environment, StrictUndefined
from markdown_it import MarkdownIt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, KeepTogether, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from game_concept.annotation import evaluate_annotations, load_annotation
from game_concept.dataset import MECHANIC_COLUMNS
from game_concept.evaluation import classification_metrics
from game_concept.mechanics import KEYS
from game_concept.predict import compare_components, concept_features, forecast
from game_concept.sources import sha256_file

REPORT = ROOT / "reports/game_concept"
OUTPUT = REPORT / "article"
TEMPLATE = ROOT / "docs/game_concept_article_template_ru.md"
MODEL_LABELS = {
    "prior": "Константный prior", "genre_only": "HGB: жанры и контекст",
    "logistic_additive": "Logistic Regression", "hgb_additive": "HGB с ограничением взаимодействий",
    "hgb_interactions": "HGB со взаимодействиями", "median": "Постоянная медиана",
    "bayesian_ridge": "Bayesian Ridge",
}
SHORT_LABELS = {**MODEL_LABELS, "hgb_additive": "HGB ограниченная", "hgb_interactions": "HGB взаимодействия"}
RU_FEATURES = {
    "log_age": "Возраст аналогов", "language_count": "Число языков", "log_price": "Цена",
    "mechanic_open_world": "Открытый мир", "mechanic_exploration": "Исследование мира",
    "genre_rpg": "Жанр RPG", "genre_casual": "Жанр Casual", "mechanic_platforming": "Платформинг",
    "mechanic_real_time_combat": "Бой в реальном времени", "mechanic_survival": "Выживание",
    "mechanic_branching_story": "Ветвящийся сюжет", "genre_simulation": "Жанр Simulation",
}
RU_MECHANICS = {
    "crafting": "Крафт", "base_building": "Строительство базы", "survival": "Выживание",
    "procedural_generation": "Процедурная генерация", "permadeath": "Перманентная смерть",
    "deck_building": "Построение колоды", "turn_based_combat": "Пошаговый бой",
    "real_time_combat": "Бой в реальном времени", "skill_tree": "Дерево навыков", "loot": "Система добычи",
    "character_customization": "Настройка персонажа", "branching_story": "Ветвящийся сюжет",
    "puzzles": "Головоломки", "platforming": "Платформинг", "stealth": "Скрытность",
    "resource_management": "Управление ресурсами", "farming": "Фермерство", "fishing": "Рыбалка",
    "taming": "Приручение", "trading": "Торговля", "automation": "Автоматизация", "physics": "Физические взаимодействия",
    "parkour": "Паркур", "rhythm": "Ритм", "time_manipulation": "Управление временем", "co_op": "Кооператив",
    "pvp": "PvP", "open_world": "Открытый мир", "user_content": "Пользовательский контент",
    "tower_defense": "Tower Defense", "exploration": "Исследование", "dating": "Отношения",
    "quests": "Квесты", "boss_battles": "Бои с боссами", "destruction": "Разрушение окружения",
    "vehicle_driving": "Управление транспортом",
}
INK, GREEN, GOLD, ROSE, GRAY = "#25282B", "#227C66", "#AD8232", "#B05670", "#7B858B"
REPRO_FILES = ["market_metrics.csv", "playtime_metrics.csv", "test_predictions.csv", "splits.csv",
               "feature_importance.csv", "temporal_backtest.csv", "calibration_metrics.csv", "reliability_bins.csv"]


def number(value, digits=4, signed=False):
    return f"{float(value):{'+' if signed else ''}.{digits}f}".replace(".", ",").replace("-", "−")


def integer(value):
    return f"{int(value):,}".replace(",", " ")


def markdown_table(headers, rows):
    def line(cells):
        return "| " + " | ".join(str(cell).replace("|", "/").replace("\n", " ") for cell in cells) + " |"
    return "\n".join([line(headers), line(["---"] * len(headers))] + [line(row) for row in rows])


def load_results(allow_retrieval_update=False):
    manifest = json.loads((REPORT / "run_manifest.json").read_text(encoding="utf-8"))
    prepared = ROOT / "data/processed/game_concept"
    audit = json.loads((prepared / "audit.json").read_text(encoding="utf-8"))
    if sha256_file(prepared / "games.csv") != manifest["input_sha256"] or audit["data_sha256"] != manifest["input_sha256"]:
        raise ValueError("Article input does not match the experiment manifest")
    changed_code = []
    for name, expected in manifest["code_sha256"].items():
        if sha256_file(ROOT / "src/game_concept" / name) != expected:
            if not allow_retrieval_update or name not in {"cli.py", "predict.py"}:
                raise ValueError(f"Experiment code changed: {name}; train and verify the new version first")
            changed_code.append(name)
    if changed_code:
        frozen = json.loads((OUTPUT / "article_manifest.json").read_text(encoding="utf-8"))
        for name in ["models/concept_model.joblib", "run_manifest.json"] + REPRO_FILES:
            key = f"reports/game_concept/{name}"
            if sha256_file(REPORT / name) != frozen["inputs_sha256"][key]:
                raise ValueError(f"Frozen research artifact changed: {name}")
    tables = {Path(name).stem: pd.read_csv(REPORT / name) for name in REPRO_FILES}
    data = pd.read_csv(prepared / "games.csv")
    splits = tables["splits"]
    if len(splits) != len(data) or splits.appid.duplicated().any() or set(splits.appid) != set(data.appid):
        raise ValueError("Saved splits do not match the prepared corpus")
    if splits.groupby("developer_group").split.nunique().gt(1).any():
        raise ValueError("Developer leakage in the saved splits")
    predictions = tables["test_predictions"]
    if set(predictions.appid) != set(splits.loc[splits.split.eq("test"), "appid"]):
        raise ValueError("Saved predictions do not match the test split")
    for row in tables["market_metrics"].itertuples(index=False):
        probs = predictions[[f"{row.model}_p{k}" for k in range(4)]].to_numpy()
        if not np.allclose(probs.sum(axis=1), 1) or (probs < 0).any():
            raise ValueError("Invalid saved probability distribution")
        metrics = classification_metrics(predictions, probs)
        if not all(np.isclose(metrics[key], getattr(row, key), atol=1e-12) for key in metrics):
            raise ValueError(f"Metrics and predictions disagree: {row.model}")
    reproduction = {name: sha256_file(REPORT / name) == sha256_file(ROOT / "reports/game_concept_repro_check" / name)
                    for name in REPRO_FILES}
    if not all(reproduction.values()):
        raise ValueError("Reproduction tables changed; verify the repeated experiment before reporting agreement")
    bundle = joblib.load(REPORT / "models/concept_model.joblib")
    if bundle["metadata"] != manifest:
        raise ValueError("Saved model and experiment manifest disagree")
    if changed_code:
        from threadpoolctl import threadpool_limits
        from game_concept.models import probabilities
        test_rows = data.set_index("appid").loc[predictions.appid]
        with threadpool_limits(limits=1):
            for name, model in bundle["all_market"].items():
                actual = probabilities(model["model"], test_rows[model["columns"]].to_numpy(dtype=float))
                expected = predictions[[f"{name}_p{i}" for i in range(4)]].to_numpy()
                if not np.allclose(actual, expected, rtol=0, atol=1e-12):
                    raise ValueError(f"Frozen model predictions changed: {name}")
    annotation_base = REPORT / "annotations"
    pointer = json.loads((annotation_base / "current.json").read_text(encoding="utf-8"))["directory"]
    if len(pointer) != 16 or any(c not in "0123456789abcdef" for c in pointer):
        raise ValueError("Invalid annotation directory")
    annotation_path = annotation_base / pointer
    tasks, _, annotation_manifest = load_annotation(annotation_path)
    if annotation_manifest["dataset_sha256"] != manifest["input_sha256"]:
        raise ValueError("Annotation audit belongs to a different corpus")
    summary, _ = evaluate_annotations(tasks, tasks[["task_id", "expert_label", "expert_notes"]])
    scrape = json.loads((prepared / "scrape_manifest.json").read_text(encoding="utf-8"))
    concept = concept_features(["Indie", "RPG"], ["crafting", "co_op"], 14.99, 3.0, 2)
    example = forecast(concept, bundle)
    if example["probabilities"] is None:
        raise ValueError("The article's demonstration concept has insufficient support")
    comparisons = compare_components(concept, bundle)
    too_many = forecast(concept_features(["Indie", "RPG"], KEYS), bundle)
    return {"manifest": manifest, "audit": audit, "data": data, "tables": tables, "reproduction": reproduction,
            "annotation": summary, "annotation_path": annotation_path, "scrape": scrape, "example": example,
            "comparisons": comparisons, "all_mechanics": too_many, "retrieval_code_changes": changed_code}


def build_context(results, author, affiliation, test_count):
    m, a, t = results["manifest"], results["audit"], results["tables"]
    market = t["market_metrics"].set_index("model")
    h = m["hypothesis"]
    chosen = market.loc[m["selected_model"]]
    calibration = t["calibration_metrics"]
    calibration = calibration[calibration.model.eq(m["selected_model"])]
    before = calibration[calibration.stage.eq("before")].set_index("threshold")
    after = calibration[calibration.stage.eq("after")].set_index("threshold")
    annotation = results["annotation"]
    temporal = t["temporal_backtest"]
    importance = t["feature_importance"].set_index("feature")
    example = results["example"]
    hours = example["playtime"]
    scenario_rows = [[f"Вероятность диапазона {label}", number(p * 100, 1) + "%"]
                     for label, p in zip(["0–20 тыс.", "20–100 тыс.", "100–500 тыс.", "500 тыс. и более"], example["probabilities"])]
    scenario_rows += [["Вероятность класса не ниже 20 тыс.", number(example["p_at_least_20k"] * 100, 1) + "%"],
                      ["Оценка медианных часов", number(hours["median_hours_estimate"], 1)],
                      ["90% интервал часов", f"{number(hours['low_hours'], 1)}–{number(hours['high_hours'], 1)}"]]
    verdicts = {
        "inconclusive": ("Убедительных свидетельств преимущества взаимодействий не получено.",
                         "Интервал пересекает ноль; преимущество взаимодействий в текущем эксперименте не подтверждено."),
        "supported": ("Получена поддержка преимущества взаимодействий в текущем эксперименте.",
                      "Интервал целиком положителен и поддерживает H1 в данном эксперименте."),
        "interactions_worse": ("Модель со взаимодействиями получила большую ошибку в текущем эксперименте.",
                               "Интервал целиком отрицателен: в данном эксперименте модель со взаимодействиями хуже."),
    }
    author_line = " | ".join(part for part in [author, affiliation] if part) or "Исследовательская статья · Game Concept Lab · 8 октября 2026 года"
    temporal_data = results["data"].sort_values(["release_date", "appid"]).iloc[int(len(results["data"]) * .8):]
    context = {
        "author_line": author_line, "sample_count": integer(a["prepared_games"]), "evidence_count": integer(a["evidence_rows"]),
        "claim_count": integer(a["feature_claims"]), "source_count": integer(a["source_rows"]), "eligible_count": integer(a["eligible_rows"]),
        "selected_label": MODEL_LABELS[m["selected_model"]], "selected_loss": number(chosen.group_log_loss),
        "prior_loss": number(market.loc["prior", "group_log_loss"]), "genre_loss": number(market.loc["genre_only", "group_log_loss"]),
        "hypothesis_delta": number(h["delta_log_loss"]), "hypothesis_low": number(h["ci_low"]),
        "hypothesis_high": number(h["ci_high"], signed=True), "hypothesis_abstract": verdicts[h["verdict"]][0],
        "hypothesis_conclusion": verdicts[h["verdict"]][1], "revision": m["snapshot_revision"],
        "taxonomy": m["taxonomy_version"], "raw_sha": a["snapshot_sha256"], "prepared_sha": m["input_sha256"],
        "hours_count": integer(a["games_with_playtime"]), "hours_share": number(a["games_with_playtime"] / len(results["data"]) * 100, 1),
        "majority_share": number(results["data"].owners_band.eq(0).mean() * 100, 1),
        "scraped_count": integer(results["scrape"]["successful"]), "html_unavailable": results["scrape"]["store_html_unavailable"],
        "connected_groups": integer(t["splits"].developer_group.nunique()), "annotation_games": annotation["games"],
        "annotation_tasks": integer(annotation["tasks"]), "annotation_reviewed": annotation["reviewed"],
        "annotation_status": "реальные precision/recall извлечения ещё не измерены независимо." if not annotation["reviewed"] else
                             "метрики рассмотренных задач не устанавливают точность всего корпуса.",
        "test_groups": integer(h["groups"]), "bootstrap_repeats": integer(h["repeats"]),
        "improvement_prior": number((1 - chosen.group_log_loss / market.loc["prior", "group_log_loss"]) * 100, 1),
        "improvement_genre": number((1 - chosen.group_log_loss / market.loc["genre_only", "group_log_loss"]) * 100, 1),
        "ece_before": number(before.loc[20000, "ece"]), "ece_after": number(after.loc[20000, "ece"]),
        "raw_loss": number(before.loc[20000, "group_log_loss"]), "calibrated_loss": number(after.loc[20000, "group_log_loss"]),
        "calibration_conclusion": "В этом эксперименте оба показателя ухудшились." if after.loc[20000, "ece"] > before.loc[20000, "ece"] and
                                 after.loc[20000, "group_log_loss"] > before.loc[20000, "group_log_loss"] else "Направление изменений зависит от показателя.",
        "temporal_train": integer(temporal.iloc[0].train_games), "temporal_test": integer(temporal.iloc[0].test_games),
        "temporal_majority": number(temporal_data.owners_band.eq(0).mean() * 100, 1),
        "playtime_selected": MODEL_LABELS[m["playtime"]["selected_model"]],
        "coverage": number(m["playtime"]["empirical_test_coverage"] * 100, 1),
        "radius": number(m["playtime"]["radius_log_hours"]), "test_hours": m["playtime"]["counts"]["test"],
        "importance_games": int(importance.evaluation_games.iloc[0]), "importance_age": number(importance.loc["log_age", "log_loss_increase"]),
        "importance_languages": number(importance.loc["language_count", "log_loss_increase"]),
        "importance_price": number(importance.loc["log_price", "log_loss_increase"]),
        "scenario_support": {"supported": "достаточная", "limited": "ограниченная", "abstain": "недостаточная"}[example["support"]["status"]],
        "density_limit": example["support"]["density_limit"], "reproduction_count": len(results["reproduction"]),
        "all_mechanics_action": "отказывается от численного прогноза" if results["all_mechanics"]["probabilities"] is None else "выдаёт численный прогноз",
        "tests_statement": f"{test_count} автоматических тестов пройдены." if test_count is not None else "число тестов в этой сборке не зафиксировано.",
        "numpy_version": m["versions"]["numpy"], "pandas_version": m["versions"]["pandas"],
        "sklearn_version": m["versions"]["scikit-learn"], "scipy_version": m["versions"]["scipy"], "lime_version": m["versions"]["lime"],
    }
    context["corpus_table"] = markdown_table(["Характеристика", "Значение"], [
        ["Строк исходного CSV", integer(a["source_rows"])], ["Игры после фильтрации", integer(a["eligible_rows"])],
        ["Подготовленная выборка", integer(a["prepared_games"])], ["Типы игровых элементов", len(MECHANIC_COLUMNS)],
        ["Записи свидетельств элементов", integer(a["evidence_rows"])], ["Кандидаты особенных возможностей", integer(a["feature_claims"])],
        ["Игры с доступными часами", integer(a["games_with_playtime"])]])
    context["split_table"] = markdown_table(["Часть", "Игры", "Группы", "Доступные часы"],
        [[name, integer(part["games"]), integer(part["developer_groups"]), integer(part["playtime_available"])] for name, part in m["splits"].items()])
    context["market_table"] = markdown_table(["Модель", "Validation L_group", "Test L_group", "Test Brier", "Balanced accuracy"],
        [[SHORT_LABELS[row.model], number(row.validation_group_log_loss), number(row.group_log_loss), number(row.brier), number(row.balanced_accuracy)]
         for row in t["market_metrics"].itertuples(index=False)])
    context["calibration_table"] = markdown_table(["Порог", "Этап", "ECE", "Binary Brier", "Group log loss"],
        [[f"{int(row.threshold / 1000)} тыс.", "До" if row.stage == "before" else "После", number(row.ece), number(row.binary_brier), number(row.group_log_loss)]
         for row in calibration.itertuples(index=False)])
    context["temporal_table"] = markdown_table(["Модель", "Group log loss", "Accuracy", "Balanced accuracy"],
        [[SHORT_LABELS[row.model], number(row.group_log_loss), number(row.accuracy), number(row.balanced_accuracy)] for row in temporal.itertuples(index=False)])
    context["playtime_table"] = markdown_table(["Модель", "Validation MAE log", "Test MAE log", "Test MAE, часы"],
        [[SHORT_LABELS[row.model], number(row.validation_mae_log_hours), number(row.test_mae_log_hours), number(row.test_mae_hours, 3)]
         for row in t["playtime_metrics"].itertuples(index=False)])
    context["scenario_table"] = markdown_table(["Показатель", "Значение"], scenario_rows)
    return context


def save_figure(fig, directory, name):
    fig.savefig(directory / name, dpi=230, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_figures(results, directory):
    directory.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.labelcolor": INK,
                         "text.color": INK, "axes.spines.top": False, "axes.spines.right": False,
                         "axes.edgecolor": "#CDD1D2", "figure.dpi": 120})
    data, tables, manifest = results["data"], results["tables"], results["manifest"]
    fig, axes = plt.subplots(1, 2, figsize=(10, 3.2), gridspec_kw={"width_ratios": [1.65, 1]})
    counts = data.owners_band.value_counts().reindex(range(4), fill_value=0)
    bars = axes[0].bar(["0–20 тыс.", "20–100 тыс.", "100–500 тыс.", "500 тыс.+"], counts, color=[GRAY, GREEN, GOLD, ROSE], width=.6)
    axes[0].bar_label(bars, labels=[integer(v) for v in counts], padding=4, fontsize=10)
    axes[0].set(ylabel="Число игр", ylim=(0, counts.max() * 1.18), title="Диапазоны владельцев, n = 8 000")
    available = int(data.playtime_median_hours.notna().sum())
    bars = axes[1].bar(["Есть оценка", "Нет оценки"], [available, len(data) - available], color=[GREEN, "#A9AFB2"], width=.55)
    axes[1].bar_label(bars, padding=4)
    axes[1].set(ylim=(0, len(data) * .9), title="Доступность положительных часов")
    fig.tight_layout(w_pad=3)
    save_figure(fig, directory, "01_dataset.png")

    fig, ax = plt.subplots(figsize=(10, 4.6))
    ax.set(xlim=(0, 10), ylim=(0, 4.6))
    ax.axis("off")
    nodes = [(0.2, 3.05, "Закреплённый CSV", "Версия, SHA-256\n125 855 исходных строк"),
             (3.55, 3.05, "Корпус и свидетельства", "Фильтры, seed = 42\n8 000 игр, 36 элементов"),
             (6.9, 3.05, "Независимые группы", "Train / validation\nCalibration / test"),
             (6.9, 1.13, "Модели и диагностика", "5 классификаторов, 3 регрессора\nBootstrap, калибровка, XAI"),
             (3.55, 1.13, "Конструктор концепции", "Вероятности, аналоги\nПоддержка и отказ от прогноза"),
             (0.2, 1.13, "Live-парсер Steam", "36 проверочных игр\nОтдельный кеш и аудит")]
    for i, (x, y, title, body) in enumerate(nodes):
        ax.add_patch(Rectangle((x, y), 2.9, 1.18, facecolor="#F3F5F4", edgecolor=GREEN if i != 5 else GRAY, linewidth=1.2))
        ax.text(x + .13, y + .92, title, fontsize=10, weight="bold", va="center")
        ax.text(x + .13, y + .48, body, fontsize=9, va="center", linespacing=1.5)
    for start, end in [((3.1, 3.64), (3.55, 3.64)), ((6.45, 3.64), (6.9, 3.64)),
                       ((8.35, 3.05), (8.35, 2.31)), ((6.9, 1.72), (6.45, 1.72))]:
        ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=14, color=GREEN))
    ax.text(1.65, .68, "Не смешивается с метками снимка", ha="center", fontsize=8.5, color=GRAY)
    save_figure(fig, directory, "02_pipeline.png")

    rows = tables["market_metrics"].copy()
    names = [SHORT_LABELS[name] for name in rows.model]
    fig, axes = plt.subplots(1, 2, figsize=(10, 3.4))
    palette = [GREEN if name == manifest["selected_model"] else GRAY for name in rows.model]
    for ax, column, title in zip(axes, ["group_log_loss", "balanced_accuracy"], ["Test group log loss ↓", "Test balanced accuracy ↑"]):
        bars = ax.barh(names, rows[column], color=palette, height=.6)
        ax.bar_label(bars, labels=[number(v) for v in rows[column]], padding=4, fontsize=9)
        ax.invert_yaxis()
        ax.set_xlim(0, rows[column].max() * 1.22)
        ax.set_title(title)
        ax.grid(axis="x", alpha=.15)
        ax.set_axisbelow(True)
    fig.tight_layout(w_pad=2.4)
    save_figure(fig, directory, "03_models.png")

    h = manifest["hypothesis"]
    fig, ax = plt.subplots(figsize=(10, 2.45))
    ax.axvline(0, color=GRAY, linestyle="--", linewidth=1)
    ax.errorbar(h["delta_log_loss"], .55, xerr=[[h["delta_log_loss"] - h["ci_low"]], [h["ci_high"] - h["delta_log_loss"]]],
                fmt="o", color=GREEN, capsize=7, linewidth=2, markersize=8)
    ax.set(xlim=(min(-.025, h["ci_low"] * 1.4), max(.025, h["ci_high"] * 1.4)), ylim=(0, 1), yticks=[],
           xlabel="Δ = L_additive − L_interactions; положительное значение означает преимущество взаимодействий")
    ax.text(h["delta_log_loss"], .79, f"Δ = {number(h['delta_log_loss'])}", ha="center", weight="bold", fontsize=11)
    ax.text(h["ci_low"], .26, number(h["ci_low"]), ha="center", fontsize=10)
    ax.text(h["ci_high"], .26, number(h["ci_high"], signed=True), ha="center", fontsize=10)
    ax.grid(axis="x", alpha=.15)
    fig.tight_layout()
    save_figure(fig, directory, "04_hypothesis.png")

    reliability = tables["reliability_bins"]
    reliability = reliability[reliability.model.eq(manifest["selected_model"])]
    fig, axes = plt.subplots(1, 3, figsize=(10, 3.35), sharex=True, sharey=True)
    for ax, threshold in zip(axes, [20000, 100000, 500000]):
        ax.plot([0, 1], [0, 1], color="#B9BFC2", linestyle="--", linewidth=1)
        for stage, label, color, marker in [("before", "До", GRAY, "s"), ("after", "После", GREEN, "o")]:
            frame = reliability[reliability.threshold.eq(threshold) & reliability.stage.eq(stage) & reliability.games.gt(0)]
            ax.plot(frame.mean_probability, frame.observed_fraction, label=label, color=color, marker=marker, markersize=3, linewidth=1.2)
            ax.scatter(frame.mean_probability, frame.observed_fraction, s=12 + 80 * frame.games / len(tables["test_predictions"]), color=color, alpha=.65)
        ax.set(title=f"Класс ≥ {threshold // 1000} тыс.", xlabel="Средняя вероятность", xlim=(-.03, 1.03), ylim=(-.03, 1.03))
        ax.grid(alpha=.15)
    axes[0].set_ylabel("Наблюдаемая доля")
    axes[-1].legend(frameon=False, fontsize=9, loc="lower right")
    fig.tight_layout(w_pad=1)
    save_figure(fig, directory, "05_calibration.png")

    importance = tables["feature_importance"].head(10).iloc[::-1]
    fig, ax = plt.subplots(figsize=(10, 4.05))
    labels = [RU_FEATURES.get(name, name) for name in importance.feature]
    ax.barh(labels, importance.log_loss_increase, xerr=importance["std"], color=[GREEN if name.startswith("mechanic_") else GRAY for name in importance.feature],
            error_kw={"elinewidth": 1, "capsize": 2, "ecolor": INK}, height=.65)
    ax.set_xlabel("Увеличение log loss при перестановке")
    ax.grid(axis="x", alpha=.15)
    ax.set_axisbelow(True)
    fig.tight_layout()
    save_figure(fig, directory, "06_importance.png")

    comparisons = results["comparisons"]
    shown = comparisons[comparisons.support_status.ne("abstain")].copy()
    shown = shown.assign(magnitude=np.maximum(shown.model_min_delta_pp.abs(), shown.model_max_delta_pp.abs())).nlargest(8, "magnitude")
    labels = [("− " if row.change == "remove" else "+ ") + RU_MECHANICS[row.mechanic] for row in shown.itertuples()]
    fig, ax = plt.subplots(figsize=(10, 3.85))
    # Recompute the individual scenario predictions rather than treating model spread as an interval.
    bundle = joblib.load(REPORT / "models/concept_model.joblib")
    base = concept_features(["Indie", "RPG"], ["crafting", "co_op"], 14.99, 3.0, 2)
    for offset, name, color, marker in [(-.17, "logistic_additive", GRAY, "s"), (0, "hgb_additive", GREEN, "o"), (.17, "hgb_interactions", ROSE, "^")]:
        baseline = forecast(base, bundle, name)["p_at_least_20k"]
        differences = []
        for row in shown.itertuples():
            variant = base.copy()
            variant.loc[0, f"mechanic_{row.mechanic}"] = 1 - int(variant.loc[0, f"mechanic_{row.mechanic}"])
            variant["mechanic_count"] = variant[MECHANIC_COLUMNS].sum(axis=1)
            differences.append((forecast(variant, bundle, name)["p_at_least_20k"] - baseline) * 100)
        ax.scatter(differences, np.arange(len(labels)) + offset, label=SHORT_LABELS[name], color=color, marker=marker, s=32)
    ax.axvline(0, linestyle="--", color="#A9AFB2", linewidth=1)
    ax.set(yticks=np.arange(len(labels)), yticklabels=labels, xlabel="Изменение вероятности, процентные пункты")
    ax.invert_yaxis()
    ax.legend(frameon=False, fontsize=8.5, loc="best")
    ax.grid(axis="x", alpha=.15)
    fig.tight_layout()
    save_figure(fig, directory, "07_scenarios.png")
    shutil.copy2(REPORT / "ui_checks/desktop.png", directory / "08_application.png")


@dataclass
class Block:
    kind: str
    html: str = ""
    level: int = 0
    rows: list[list[str]] = field(default_factory=list)
    path: Path | None = None
    caption: str = ""


def parse_markdown(markdown, output):
    renderer = MarkdownIt("commonmark").enable("table")
    soup = BeautifulSoup(renderer.render(markdown), "html.parser")
    blocks = []
    for element in soup.children:
        if not isinstance(element, Tag):
            continue
        if element.name in ["h1", "h2", "h3", "h4"]:
            blocks.append(Block("heading", element.decode_contents(), int(element.name[1])))
        elif element.name == "p":
            image = element.find("img")
            if image:
                path = (output / image["src"]).resolve()
                if not path.is_relative_to(output.resolve()) or not path.exists():
                    raise ValueError("Article image is missing or outside its output directory")
                blocks.append(Block("figure", path=path, caption=image.get("alt", "")))
            else:
                kind = "caption" if element.get_text().startswith("Таблица ") else "paragraph"
                blocks.append(Block(kind, element.decode_contents()))
        elif element.name == "table":
            rows = [[cell.decode_contents() for cell in row.find_all(["th", "td"], recursive=False)] for row in element.find_all("tr")]
            if not rows or any(len(row) != len(rows[0]) for row in rows):
                raise ValueError("Invalid article table")
            blocks.append(Block("table", rows=rows))
        elif element.name in ["ol", "ul"]:
            start = int(element.get("start", 1))
            for index, item in enumerate(element.find_all("li", recursive=False), start):
                prefix = f"{index}. " if element.name == "ol" else "• "
                blocks.append(Block("list", prefix + item.decode_contents()))
        elif element.name == "pre":
            blocks.append(Block("code", element.get_text().rstrip()))
        else:
            raise ValueError(f"Unsupported article block: {element.name}")
    return blocks


def plain(html):
    return BeautifulSoup(html, "html.parser").get_text()


def append_inline(paragraph, html):
    def add(nodes, bold=False, italic=False, mono=False, link=None):
        for node in nodes:
            if isinstance(node, NavigableString):
                run = paragraph.add_run(str(node))
                run.bold, run.italic = bold, italic
                if mono:
                    run.font.name, run.font.size = "Consolas", Pt(9)
                if link:
                    relation = paragraph.part.relate_to(link, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
                    hyperlink = OxmlElement("w:hyperlink")
                    hyperlink.set(qn("r:id"), relation)
                    run.font.color.rgb = RGBColor.from_string("227C66")
                    hyperlink.append(run._r)
                    paragraph._p.append(hyperlink)
            elif node.name == "br":
                paragraph.add_run().add_break()
            else:
                add(node.children, bold or node.name in ["b", "strong"], italic or node.name in ["i", "em"],
                    mono or node.name == "code", node.get("href") if node.name == "a" else link)
    add(BeautifulSoup(html, "html.parser").children)


def image_dimensions(path, max_width, max_height):
    from PIL import Image as PILImage
    with PILImage.open(path) as picture:
        width, height = picture.size
    ratio = min(max_width / width, max_height / height)
    return width * ratio, height * ratio


def write_docx(blocks, path, author):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.1)
    section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [("Title", 17), ("Heading 1", 13), ("Heading 2", 11.5)]:
        style = document.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor.from_string("25282B")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
    document.styles["Caption"].font.name = "Times New Roman"
    document.styles["Caption"].font.size = Pt(9)
    header = section.header.paragraphs[0]
    header.text = "GAME CONCEPT LAB  |  ИССЛЕДОВАТЕЛЬСКАЯ СТАТЬЯ"
    header.runs[0].font.name, header.runs[0].font.size = "Arial", Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("7B858B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("8 октября 2026  |  ").font.size = Pt(8)
    field_node = OxmlElement("w:fldSimple")
    field_node.set(qn("w:instr"), "PAGE")
    footer._p.append(field_node)
    document.core_properties.title = plain(next(b.html for b in blocks if b.kind == "heading"))
    document.core_properties.author = author
    document.core_properties.subject = "Steam: ретроспективная оценка аудитории по жанру и игровым элементам"
    for block in blocks:
        if block.kind == "heading":
            style = "Title" if block.level == 1 else f"Heading {block.level - 1}"
            paragraph = document.add_paragraph(style=style)
            append_inline(paragraph, block.html)
        elif block.kind in ["paragraph", "list", "caption"]:
            paragraph = document.add_paragraph(style="Caption" if block.kind == "caption" else "Normal")
            append_inline(paragraph, block.html)
            if block.kind == "caption":
                paragraph.paragraph_format.keep_with_next = True
            elif block.kind == "paragraph":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.html)
            run.font.name, run.font.size = "Consolas", Pt(8.5)
            paragraph.paragraph_format.keep_together = True
        elif block.kind == "figure":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            width, _ = image_dimensions(block.path, 17, 12.8)
            paragraph.add_run().add_picture(str(block.path), width=Cm(width))
            caption = document.add_paragraph(block.caption, style="Caption")
            caption.paragraph_format.keep_together = True
        elif block.kind == "table":
            table = document.add_table(rows=0, cols=len(block.rows[0]))
            table.style = "Table Grid"
            for index, cells in enumerate(block.rows):
                row = table.add_row()
                properties = row._tr.get_or_add_trPr()
                properties.append(OxmlElement("w:cantSplit"))
                if index == 0:
                    properties.append(OxmlElement("w:tblHeader"))
                for cell, value in zip(row.cells, cells):
                    append_inline(cell.paragraphs[0], value)
                    cell.paragraphs[0].paragraph_format.space_after = Pt(3)
                    cell.paragraphs[0].paragraph_format.line_spacing = 1
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9)
                        run.bold = index == 0
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EAF0ED" if index == 0 else "FFFFFF")
                    cell._tc.get_or_add_tcPr().append(shade)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
    document.save(path)


def register_fonts():
    families = [
        (Path("C:/Windows/Fonts"), ["times.ttf", "timesbd.ttf", "timesi.ttf", "timesbi.ttf"]),
        (Path("/usr/share/fonts/truetype/liberation2"), ["LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf", "LiberationSerif-Italic.ttf", "LiberationSerif-BoldItalic.ttf"]),
        (Path("/usr/share/fonts/truetype/dejavu"), ["DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf", "DejaVuSerif-Italic.ttf", "DejaVuSerif-BoldItalic.ttf"]),
    ]
    for base, names in families:
        if all((base / name).exists() for name in names):
            for font, name in zip(["Article", "ArticleBold", "ArticleItalic", "ArticleBoldItalic"], names):
                pdfmetrics.registerFont(TTFont(font, str(base / name)))
            pdfmetrics.registerFontFamily("Article", normal="Article", bold="ArticleBold", italic="ArticleItalic", boldItalic="ArticleBoldItalic")
            break
    else:
        raise FileNotFoundError("Install Times New Roman, Liberation Serif or DejaVu Serif with Cyrillic support")


def pdf_inline(html):
    soup = BeautifulSoup(html, "html.parser")
    for node in soup.find_all("strong"):
        node.name = "b"
    for node in soup.find_all("em"):
        node.name = "i"
    for node in soup.find_all("code"):
        node.name, node.attrs = "font", {"name": "Article", "size": "9"}
    for node in soup.find_all("a"):
        node.name, node.attrs = "link", {"href": node.get("href"), "color": GREEN}
    for node in soup.find_all("p"):
        node.unwrap()
    return str(soup)


def write_pdf(blocks, path, author):
    register_fonts()
    width = A4[0] - 40 * mm
    styles = {
        "body": ParagraphStyle("Body", fontName="Article", fontSize=11, leading=14.2, spaceAfter=7, alignment=TA_JUSTIFY,
                               allowWidows=0, allowOrphans=0, splitLongWords=1),
        "list": ParagraphStyle("List", fontName="Article", fontSize=10.5, leading=13.7, spaceAfter=5, splitLongWords=1),
        "title": ParagraphStyle("ArticleTitle", fontName="ArticleBold", fontSize=17, leading=20.5, spaceAfter=12, keepWithNext=True),
        "heading": ParagraphStyle("Section", fontName="ArticleBold", fontSize=13, leading=16, spaceBefore=12, spaceAfter=7, keepWithNext=True),
        "subheading": ParagraphStyle("Subsection", fontName="ArticleBold", fontSize=11.5, leading=14, spaceBefore=8, spaceAfter=6, keepWithNext=True),
        "caption": ParagraphStyle("Caption", fontName="ArticleItalic", fontSize=9, leading=11.3, spaceAfter=7),
        "table": ParagraphStyle("Cell", fontName="Article", fontSize=9, leading=11, splitLongWords=1),
        "code": ParagraphStyle("Code", fontName="Courier", fontSize=8.2, leading=11, spaceAfter=8),
    }
    flow = []
    for block in blocks:
        if block.kind == "heading":
            style = styles["title" if block.level == 1 else "heading" if block.level == 2 else "subheading"]
            flow.append(Paragraph(pdf_inline(block.html), style))
        elif block.kind in ["paragraph", "list", "caption"]:
            flow.append(Paragraph(pdf_inline(block.html), styles["body" if block.kind == "paragraph" else block.kind]))
            if block.kind == "caption":
                flow[-1].keepWithNext = True
        elif block.kind == "code":
            flow.append(Preformatted(block.html, styles["code"]))
        elif block.kind == "figure":
            image_width, image_height = image_dimensions(block.path, width, 350)
            picture = Image(str(block.path), width=image_width, height=image_height)
            picture.hAlign = "CENTER"
            flow.append(KeepTogether([Spacer(1, 5), picture, Spacer(1, 4), Paragraph(block.caption, styles["caption"])]))
        elif block.kind == "table":
            count = len(block.rows[0])
            proportions = [.62, .38] if count == 2 else [.32] + [.68 / (count - 1)] * (count - 1)
            cells = [[Paragraph(pdf_inline(value), styles["table"]) for value in row] for row in block.rows]
            table = Table(cells, colWidths=[width * share for share in proportions], repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF0ED")),
                ("LINEABOVE", (0, 0), (-1, 0), .6, colors.HexColor(GREEN)),
                ("LINEBELOW", (0, 0), (-1, 0), .4, colors.HexColor(GRAY)),
                ("LINEBELOW", (0, -1), (-1, -1), .5, colors.HexColor(GRAY)),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F7F6")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            flow.extend([table, Spacer(1, 8)])
    def page_frame(canvas, document):
        canvas.saveState()
        canvas.setFont("Article", 8)
        canvas.setFillColor(colors.HexColor(GRAY))
        canvas.drawString(20 * mm, A4[1] - 13 * mm, "GAME CONCEPT LAB  |  ИССЛЕДОВАТЕЛЬСКАЯ СТАТЬЯ")
        canvas.setStrokeColor(colors.HexColor("#D4D9D6"))
        canvas.line(20 * mm, A4[1] - 15 * mm, A4[0] - 20 * mm, A4[1] - 15 * mm)
        canvas.drawString(20 * mm, 12 * mm, "8 октября 2026 года")
        canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, str(document.page))
        canvas.restoreState()
    title = plain(next(b.html for b in blocks if b.kind == "heading"))
    document = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=22 * mm,
                                 bottomMargin=20 * mm, title=title, author=author, invariant=1, pageCompression=1)
    document.build(flow, onFirstPage=page_frame, onLaterPages=page_frame)


def main():
    parser = argparse.ArgumentParser(description="Build the Steam concept research article without changing model results")
    parser.add_argument("--author", default="")
    parser.add_argument("--affiliation", default="")
    parser.add_argument("--verified-tests", type=int, help="Number of tests independently run successfully for this release")
    args = parser.parse_args()
    results = load_results()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    make_figures(results, OUTPUT / "figures")
    context = build_context(results, args.author, args.affiliation, args.verified_tests)
    environment = Environment(undefined=StrictUndefined, autoescape=False, keep_trailing_newline=True)
    article = environment.from_string(TEMPLATE.read_text(encoding="utf-8")).render(**context)
    source = OUTPUT / "article_ru.md"
    source.write_text(article, encoding="utf-8")
    blocks = parse_markdown(article, OUTPUT)
    write_docx(blocks, OUTPUT / "article_ru.docx", args.author)
    write_pdf(blocks, OUTPUT / "article_ru.pdf", args.author)
    results["comparisons"].to_csv(OUTPUT / "scenario_comparison.csv", index=False)
    demonstration = {key: value for key, value in results["example"].items() if key != "analogues"}
    (OUTPUT / "demonstration.json").write_text(json.dumps(demonstration, ensure_ascii=False, indent=2), encoding="utf-8")
    inputs = [TEMPLATE, ROOT / "requirements-article.txt", Path(__file__), ROOT / "data/processed/game_concept/audit.json",
              ROOT / "data/processed/game_concept/scrape_manifest.json", REPORT / "run_manifest.json",
              REPORT / "models/concept_model.joblib", REPORT / "ui_checks/desktop.png"]
    inputs += [REPORT / name for name in REPRO_FILES]
    inputs += [results["annotation_path"] / name for name in ["manifest.json", "tasks.csv", "games.json"]]
    if (results["annotation_path"] / "reviews.csv").exists():
        inputs.append(results["annotation_path"] / "reviews.csv")
    manifest = {"article_date": "2026-10-08", "author": args.author, "affiliation": args.affiliation,
                "experiment_input_sha256": results["manifest"]["input_sha256"], "tests_reported": args.verified_tests,
                "reproduction": results["reproduction"], "figures": sum(b.kind == "figure" for b in blocks),
                "tables": sum(b.kind == "table" for b in blocks), "annotation": results["annotation"],
                "inputs_sha256": {path.relative_to(ROOT).as_posix(): sha256_file(path) for path in inputs},
                "outputs_sha256": {path.relative_to(OUTPUT).as_posix(): sha256_file(path) for path in sorted(OUTPUT.glob("article_ru.*"))},
                "image_sha256": {path.name: sha256_file(path) for path in sorted((OUTPUT / "figures").glob("*.png"))},
                "build_versions": {name: importlib.metadata.version(name) for name in ["python-docx", "reportlab", "markdown-it-py", "Jinja2", "matplotlib"]},
                "status": "complete article text; human annotation and prospective validation unfinished"}
    (OUTPUT / "article_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Article created: {OUTPUT}")
    print(f"Figures: {manifest['figures']}; tables: {manifest['tables']}; scenario: {demonstration['support']['status']}")


if __name__ == "__main__":
    main()
