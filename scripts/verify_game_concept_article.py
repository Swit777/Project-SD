from __future__ import annotations

import json
from pathlib import Path
import re
import sys

from docx import Document
from markdown_it import MarkdownIt
import numpy as np
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from game_concept.sources import sha256_file

ARTICLE = ROOT / "reports/game_concept/article"


def verify():
    manifest = json.loads((ARTICLE / "article_manifest.json").read_text(encoding="utf-8"))
    for name, expected in manifest["outputs_sha256"].items():
        if sha256_file(ARTICLE / name) != expected:
            raise AssertionError(f"Article output changed: {name}")
    for name, expected in manifest["inputs_sha256"].items():
        if sha256_file(ROOT / name) != expected:
            raise AssertionError(f"Article source changed: {name}; rebuild")
    for name, expected in manifest["image_sha256"].items():
        if sha256_file(ARTICLE / "figures" / name) != expected:
            raise AssertionError(f"Article image changed: {name}")
    markdown = (ARTICLE / "article_ru.md").read_text(encoding="utf-8")
    if any(marker in markdown for marker in ["{{", "}}", "\ufffd", "???"]):
        raise AssertionError("Unrendered template or corrupted Cyrillic")
    tokens = MarkdownIt("commonmark").enable("table").parse(markdown)
    figures = [child for token in tokens for child in (token.children or []) if child.type == "image"]
    assert len(figures) == manifest["figures"] == 8
    assert sum(token.type == "table_open" for token in tokens) == manifest["tables"] == 7
    docx = Document(ARTICLE / "article_ru.docx")
    assert len(docx.inline_shapes) == manifest["figures"]
    assert len(docx.tables) == manifest["tables"]
    pdf = pymupdf.open(ARTICLE / "article_ru.pdf")
    full_text = "\n".join(page.get_text() for page in pdf)
    docx_text = "\n".join(paragraph.text for paragraph in docx.paragraphs)
    for label in ["Аннотация", "Введение", "Заключение", "Список источников", "Приложение А"]:
        assert label in full_text and label in docx_text, label
    for index in range(1, manifest["figures"] + 1):
        assert f"Рисунок {index}." in full_text and f"Рисунок {index}." in docx_text
    for index in range(1, manifest["tables"] + 1):
        assert f"Таблица {index}." in full_text and f"Таблица {index}." in docx_text
    assert "\ufffd" not in full_text and "???" not in full_text
    references = markdown.split("## Список источников\n", 1)[1].split("## Приложение", 1)[0]
    assert len(re.findall(r"(?m)^\d+\. ", references)) == 13
    table_text = "\n".join(cell.text for table in docx.tables for row in table.rows for cell in row.cells)
    for value in ["0,5954", "0,7104", "0,6443", "0,0271", "5,851"]:
        assert value in full_text and value in table_text, value
    assert "не подтверждено" in full_text
    out = ARTICLE / "verification"
    out.mkdir(exist_ok=True)
    pages = []
    for page in pdf:
        text = page.get_text()
        if len(text) < 300:
            raise AssertionError(f"Empty or almost empty page {page.number + 1}")
        lines = []
        for block in page.get_text("dict")["blocks"]:
            if block["type"] != 0:
                continue
            for line in block["lines"]:
                rect = pymupdf.Rect(line["bbox"])
                if rect.x0 < 50 or rect.x1 > page.rect.width - 50 or rect.y0 < 20 or rect.y1 > page.rect.height - 20:
                    raise AssertionError(f"Text exceeds page bounds: page {page.number + 1}, {rect}")
                lines.append(rect)
        for index, rect in enumerate(lines):
            for other in lines[index + 1:]:
                overlap = rect & other
                if not overlap.is_empty and overlap.width > 3 and overlap.height > 2:
                    raise AssertionError(f"Overlapping text lines on page {page.number + 1}")
        picture = page.get_pixmap(matrix=pymupdf.Matrix(1.2, 1.2), alpha=False)
        pixels = np.frombuffer(picture.samples, dtype=np.uint8).reshape(picture.height, picture.width, picture.n)
        ink_fraction = float(np.mean(pixels[:, :, :3].min(axis=2) < 230))
        if ink_fraction < .01:
            raise AssertionError(f"Blank rendered page {page.number + 1}")
        picture.save(out / f"page_{page.number + 1:02d}.png")
        images = page.get_images(full=True)
        for image in images:
            for rect in page.get_image_rects(image[0]):
                if not page.rect.contains(rect) or rect.is_empty:
                    raise AssertionError("Missing or clipped embedded figure")
        pages.append({"page": page.number + 1, "characters": len(text), "images": len(images), "ink_fraction": ink_fraction})
    assert sum(page["images"] for page in pages) == manifest["figures"]
    checks = {"status": "passed", "pages": len(pdf), "figures": manifest["figures"], "tables": manifest["tables"],
              "references": 13, "cyrillic_text": "verified", "hashes": "verified", "page_bounds": "verified",
              "text_overlaps": "none detected", "rendered_pages": pages}
    (out / "checks.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    pdf.close()
    print(json.dumps({key: value for key, value in checks.items() if key != "rendered_pages"}, ensure_ascii=False))


if __name__ == "__main__":
    verify()
