from __future__ import annotations

import json
from pathlib import Path
import re
import sys

from docx import Document
from markdown_it import MarkdownIt
import numpy as np
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from game_concept.sources import sha256_file

ARTICLE = ROOT / "reports/game_concept/article_by_example"
MAX_PAGES = 10


def verify():
    manifest = json.loads((ARTICLE / "manifest.json").read_text(encoding="utf-8"))
    docx_path = ROOT / manifest["output_docx"]
    checked = {
        docx_path: manifest["docx_sha256"],
        ROOT / manifest["formatting_reference"]: manifest["reference_sha256"],
        ARTICLE / "article_body_ru.md": manifest["article_body_sha256"],
        ROOT / "reports/game_concept/analogue_upgrade/comparison.json": manifest["retrieval_comparison_sha256"],
        ROOT / "reports/game_concept/analogue_upgrade/tests.log": manifest["test_log_sha256"],
    }
    checked.update({ARTICLE / "figures" / name: digest for name, digest in manifest["figures_sha256"].items()})
    checked.update({ROOT / ("docs" if name.endswith(".md") else "scripts") / name: digest
                    for name, digest in manifest["source_code_sha256"].items()})
    for path, expected in checked.items():
        if sha256_file(path) != expected:
            raise AssertionError(f"Article source or output changed: {path}")

    markdown = (ARTICLE / "article_body_ru.md").read_text(encoding="utf-8")
    assert not any(marker in markdown for marker in ["{{", "}}", "\ufffd", "???"])
    tokens = MarkdownIt("commonmark").enable("table").parse(markdown)
    figures = [child for token in tokens for child in (token.children or []) if child.type == "image"]
    document = Document(docx_path)
    assert len(figures) == len(document.inline_shapes) == manifest["figures"] == 5
    assert sum(token.type == "table_open" for token in tokens) == len(document.tables) == manifest["tables"] == 2
    body, bibliography = markdown.split("## Библиография\n", 1)
    references = {int(value) for value in re.findall(r"(?m)^\[(\d+)\]\t", bibliography)}
    citations = {int(value) for value in re.findall(r"\b\d+\b", " ".join(re.findall(r"\[([\d, ]+)\]", body)))}
    assert references == citations == set(range(1, manifest["sources"] + 1))

    pdf_path = ARTICLE / "article_ru.pdf"
    with pymupdf.open(pdf_path) as pdf:
        assert 1 <= len(pdf) <= MAX_PAGES, f"The article has {len(pdf)} pages; maximum {MAX_PAGES}"
        text = "\n".join(page.get_text() for page in pdf)
        docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for label in ["Резюме", "Введение", "Данные и методология", "Результаты и обсуждение", "Выводы", "Библиография"]:
            assert label in text and label in docx_text, label
        for label, count in [("Рисунок", manifest["figures"]), ("Таблица", manifest["tables"])]:
            for index in range(1, count + 1):
                assert f"{label} {index}" in text and f"{label} {index}" in docx_text
        for value in ["0,5954", "0,7104", "0,5989", "94 016", "не подтверждено", "100 тестов"]:
            assert value in text, value
        assert "\ufffd" not in text and "???" not in text

        out = ARTICLE / "verification_short_2026_10_09"
        out.mkdir(exist_ok=True)
        pages = []
        for page in pdf:
            lines = []
            for block in page.get_text("dict")["blocks"]:
                if block["type"] != 0:
                    continue
                for line in block["lines"]:
                    rect = pymupdf.Rect(line["bbox"])
                    if rect.x0 < 40 or rect.x1 > page.rect.width - 30 or rect.y0 < 15 or rect.y1 > page.rect.height - 15:
                        raise AssertionError(f"Text exceeds page bounds: page {page.number + 1}, {rect}")
                    lines.append(rect)
            for index, rect in enumerate(lines):
                for other in lines[index + 1:]:
                    overlap = rect & other
                    if not overlap.is_empty and overlap.width > 3 and overlap.height > 2:
                        raise AssertionError(f"Overlapping text on page {page.number + 1}")
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.3, 1.3), alpha=False)
            pixels = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(pixmap.height, pixmap.width, pixmap.n)
            ink_fraction = float(np.mean(pixels[:, :, :3].min(axis=2) < 230))
            assert len(page.get_text()) > 300 and ink_fraction > .01, f"Almost empty page {page.number + 1}"
            for entry in page.get_images(full=True):
                for rect in page.get_image_rects(entry[0]):
                    assert page.rect.contains(rect) and not rect.is_empty, "Clipped figure"
            pixmap.save(out / f"page_{page.number + 1:02d}.png")
            pages.append({"page": page.number + 1, "characters": len(page.get_text()), "ink_fraction": ink_fraction})
        checks = {
            "status": "passed", "pages": len(pdf), "max_pages": MAX_PAGES,
            "includes_cover_and_bibliography": True, "figures": manifest["figures"], "tables": manifest["tables"],
            "references": manifest["sources"], "reviewed_tests": manifest["reviewed_tests"],
            "hashes": "verified", "cyrillic_text": "verified", "page_bounds": "verified",
            "text_overlaps": "none detected", "docx_sha256": sha256_file(docx_path),
            "pdf_sha256": sha256_file(pdf_path), "rendered_pages": pages,
        }
    (out / "checks.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in checks.items() if key != "rendered_pages"}, ensure_ascii=False))


if __name__ == "__main__":
    verify()
